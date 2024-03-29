from openpyxl.workbook import Workbook
from openpyxl import load_workbook
from openpyxl.chart import PieChart, Reference, Series, PieChart3D

wb = Workbook()
# Get sheet
ws = wb.active


data = [
    ['Flavor', 'Sold'],
    ['Vanilla', 1500],
    ['Chocolate', 1700],
    ['Strawberry', 600],
    ['Pumpkin', 950]
]

# add data to wook sheet
for rows in data:
  ws.append(rows)

chart = PieChart()
labels = Reference(ws, min_col=1, min_row=2, max_row=5)
data = Reference(ws, min_col=2, min_row=2, max_row=5)

chart.add_data(data, titles_from_data=True)
chart.set_categories(labels)
chart.title = 'Ice Cream by Flavor'

ws.add_chart(chart, 'C1')
wb.save('target/Pie.xlsx')