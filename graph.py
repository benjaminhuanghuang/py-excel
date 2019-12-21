from openpyxl.workbook import Workbook
from openpyxl import load_workbook
from openpyxl.chart import PieChart, Reference, Series, BarChart

wb = load_workbook('data/crime_report.xlsx')
# Get sheet
ws = wb.active


chart = BarChart()

data = Reference(ws, min_row=8, min_col=1, max_col=13, max_row=13)
labels = Reference(ws, min_row=8, min_col=2, max_col=8, max_row=13)

chart.add_data(data, from_rows=True, titles_from_data=True)
chart.set_categories(labels)
chart.title = 'Crimes'
chart.height = 10
chart.width = 20
ws.add_chart(chart, 'B14')

wb.save('target/lines.xlsx')
