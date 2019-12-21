from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
from openpyxl.workbook import Workbook
from openpyxl import load_workbook

wb = load_workbook('data/Pie.xlsx')
ws = wb.active

tab = Table(displayName='Table1', ref='A1:B5')

style = TableStyleInfo(name='TableStyleMedium9', 
showFirstColumn =False,
showLastColumn=False,
showRowStripes=True,
showColumnStripes=True)

tab.tabelStyle = style
ws.add_table(tab)


#  Add image
img = Image('data/madecraft.jpg')
img.height = img.height * .25
img.width = img.width * .25

ws.add_image(img, 'C1')
wb.save('target/table.xlsx')


